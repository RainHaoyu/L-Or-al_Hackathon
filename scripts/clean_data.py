"""数据层清洗脚本：xlsx → data/seed/{allergens,ige_materials}.json

可重复执行（数据热更新工作流的一部分）：
    python scripts/clean_data.py [--src <数据层目录>] [--out <seed目录>]
"""
from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SRC = REPO_ROOT.parent / "数据层"
DEFAULT_OUT = REPO_ROOT / "data" / "seed"

# 高频致敏香料定性备注（来自《香水成分数据及过敏香料.docx》，按 INCI 名合并）
HIGH_FREQUENCY_NOTES: dict[str, str] = {
    "d-Limonene": "柑橘精油来源，香柠檬/橙子/柑橘调；氧化后致敏性大幅上升，久放香水风险变高",
    "Linalool": "薰衣草、佛手柑、橙花；氧化产物刺激性很强，最常见香料过敏原之一",
    "Geraniol": "玫瑰、天竺葵，甜花香；花香调香水几乎普遍含有",
    "Citral": "强烈柠檬香气，柑橘、花香调；接触性皮炎高发",
    "Eugenol": "丁香油，辛香、东方调；刺激性较强，IFRA 严格限用量",
    "Cinnamal": "肉桂香气，辛香调，强致敏物",
    "Coumarin": "香草、干草甜香，美食调；兼具光敏性，晒太阳更容易泛红发痒",
    "Citronellol": "玫瑰花香原料，广泛用于花香香水",
    "Isoeugenol": "温暖辛香，IFRA 极低上限，高致敏性",
    "Hydroxycitronellal": "经典铃兰花香原料，花香调高频成分",
    "Farnesol": "茉莉、橙花，花香柑橘调；易氧化致敏",
    "Benzyl salicylate": "花香定香剂，很多白花调香水都有",
    "Oakmoss extract": "西普调经典原料，致敏极强，IFRA 大幅限制使用量",
}

# 文档算例给出的唯一确定性 NESIL（柠檬烯，NESIL=10000, SAF=100）
DOCUMENTED_NESIL = {"d-Limonene": 10000.0}
# 无公开数据成分的演示用估计值（tier→NESIL），引擎会叠加 demo_estimate_penalty 并标注 indicative
DEMO_ESTIMATE_NESIL = {"strong": 1000.0, "medium": 5000.0, "low": 10000.0}
DOCUMENTED_K_OX = {"d-Limonene": 0.03, "Linalool": 0.0175}

# 备注→香调家族关键词映射（与 families.json 的 7 个 key 对齐）
FAMILY_KEYWORDS: list[tuple[str, str]] = [
    ("茉莉", "floral"), ("玫瑰", "floral"), ("花香", "floral"), ("白花", "floral"), ("铃兰", "floral"),
    ("天竺葵", "floral"), ("橙花", "floral"), ("晚香玉", "floral"), ("康乃馨", "floral"),
    ("肉桂", "oriental"), ("丁香", "oriental"), ("辛香", "oriental"), ("香膏", "oriental"), ("树脂", "oriental"),
    ("柑橘", "citrus"), ("柠檬", "citrus"), ("橙", "citrus"), ("果皮", "citrus"), ("柠檬草", "citrus"),
    ("苔", "fougere"), ("西普", "fougere"), ("馥奇", "fougere"), ("薰衣草", "fougere"),
    ("紫罗兰叶", "aquatic"), ("水生", "aquatic"), ("青香", "aquatic"),
    ("香草", "gourmand"), ("美食", "gourmand"), ("蜂蜜", "gourmand"), ("甜香", "gourmand"),
    ("木质", "woody"), ("定香", "woody"),
]


def norm(s: str | None) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s))
    for h in "\u2010\u2011\u2012\u2013\u2014\u2212":  # 各类 Unicode 连字符 → ASCII
        s = s.replace(h, "-")
    return re.sub(r"\s+", " ", s).strip()


def norm_cas(s: str) -> str:
    return norm(s).replace("‑", "-").replace("–", "-").replace("—", "-")


def classify_tier(note: str) -> str:
    if any(k in note for k in ("强", "极强", "严格限用", "大幅限制")):
        return "strong"
    if any(k in note for k in ("氧化", "升高", "高发", "高频")):
        return "medium"
    return "low"


def infer_families(note: str) -> list[str]:
    found: list[str] = []
    for kw, fam in FAMILY_KEYWORDS:
        if kw in note and fam not in found:
            found.append(fam)
    return found or ["floral"]


def parse_allergens(src_dir: Path) -> list[dict]:
    wb = openpyxl.load_workbook(src_dir / "表格26种致敏香料.xlsx", read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    out: list[dict] = []
    for r in rows:
        cells = [norm(c) for c in (list(r) + [None] * 5)[:5]]
        if not cells[0].isdigit():
            continue
        _, inci, name_zh, cas, note = cells
        inci_c = inci.split(" (")[0].strip()
        tier = classify_tier(note)
        high_freq = inci_c in HIGH_FREQUENCY_NOTES
        oxidation_prone = ("氧化" in note) or (inci_c in DOCUMENTED_K_OX)
        if inci_c in DOCUMENTED_NESIL:
            nesil, nesil_source = DOCUMENTED_NESIL[inci_c], "documented"
        else:
            nesil, nesil_source = DEMO_ESTIMATE_NESIL[tier], "demo_estimate"
        out.append({
            "inci": inci_c,
            "inci_raw": inci,
            "name_zh": name_zh,
            "cas": norm_cas(cas),
            "tier": tier,
            "high_frequency": high_freq,
            "oxidation_prone": oxidation_prone,
            "k_ox_per_day": DOCUMENTED_K_OX.get(inci_c),
            "banned_eu": "禁用" in note,
            "reproductive_flag": "Lilial" in inci_c or "Butylphenyl" in inci_c,
            "families": infer_families(note + " " + HIGH_FREQUENCY_NOTES.get(inci_c, "")),
            "nesil": nesil,
            "nesil_source": nesil_source,
            "ifra_limit_pct": None,
            "clinical_noel": None,
            "note": note,
        })
    return out


IGE_LEVEL_MAP = {"高": "high", "中": "medium", "低": "low"}


def parse_ige_level(s: str) -> str:
    if "中" in s and "低" not in s:
        return "medium"
    if "中‑低" in s or "中-低" in s:
        return "medium_low"
    if s.startswith("低"):
        return "low"
    return "medium"


def parse_ige(src_dir: Path) -> list[dict]:
    wb = openpyxl.load_workbook(src_dir / "IgE致敏原表格.xlsx", read_only=True, data_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    wb.close()
    out: list[dict] = []
    header_seen = False
    for r in rows:
        cells = [norm(c) for c in (list(r) + [None] * 7)[:7]]
        non_empty = [c for c in cells if c]
        if not non_empty:
            continue
        if cells[2] in ("英文名 / INCI", "原料类型") or cells[1] == "原料中文名":
            header_seen = True
            continue
        if not header_seen or len(cells[1]) < 2 or cells[1] == "原料中文名":
            continue
        # 树脂/净油节：原料中文名 | 英文名/INCI | 来源类型 | 常见香调 | IgE风险 | 备注
        if cells[2] and cells[3]:
            out.append({
                "name_zh": cells[1],
                "inci": cells[2],
                "source_type": cells[3],
                "families": infer_families(cells[4]),
                "families_zh": cells[4],
                "ige_risk": parse_ige_level(cells[5]),
                "note": cells[6],
            })
        elif cells[2] and "诱发机制" not in cells[2]:
            # 花粉节（原料类型 | 代表原料 | 机制 | 表现）
            out.append({
                "name_zh": cells[1],
                "inci": "",
                "source_type": cells[2],
                "families": infer_families(cells[3] or ""),
                "families_zh": cells[3] or "",
                "ige_risk": "medium",
                "note": f"{cells[4] or ''}；{cells[5] or ''}",
            })
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", type=Path, default=DEFAULT_SRC)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = ap.parse_args()

    allergens = parse_allergens(args.src)
    ige = parse_ige(args.src)

    args.out.mkdir(parents=True, exist_ok=True)
    # 防覆盖守卫：allergens v2（build_allergens_v2.py 产物，含文献级 NESIL/EU 45 新条目）
    # 存在时，本脚本的基础产物改写到 *.base.json，避免清掉增强数据
    existing = args.out / "allergens.json"
    v2_present = False
    if existing.exists():
        try:
            n = len(json.loads(existing.read_text(encoding="utf-8"))["items"])
            v2_present = n > len(allergens)
        except Exception:
            v2_present = False
    out_name = "allergens.base.json" if v2_present else "allergens.json"
    if v2_present:
        print(f"[guard] 检测到 v2 增强数据（{n} 条 > 基础 {len(allergens)} 条），"
              f"基础产物写入 {out_name}；如需重建 v2 请运行 scripts/build_allergens_v2.py")

    (args.out / out_name).write_text(
        json.dumps({"$comment": "由 scripts/clean_data.py 从《表格26种致敏香料.xlsx》生成并合并 docx 定性备注", "items": allergens},
                   ensure_ascii=False, indent=2), encoding="utf-8")
    (args.out / "ige_materials.json").write_text(
        json.dumps({"$comment": "由 scripts/clean_data.py 从《IgE致敏原表格.xlsx》生成（IgE Ⅰ 型速发风险，用于鼻炎人群提示）", "items": ige},
                   ensure_ascii=False, indent=2), encoding="utf-8")

    strong = [a["inci"] for a in allergens if a["tier"] == "strong"]
    ox = [a["inci"] for a in allergens if a["oxidation_prone"]]
    banned = [a["inci"] for a in allergens if a["banned_eu"]]
    documented = [a["inci"] for a in allergens if a["nesil_source"] == "documented"]
    print(f"allergens: {len(allergens)} 条 | strong: {strong}")
    print(f"oxidation_prone: {ox}")
    print(f"banned_eu: {banned}")
    print(f"nesil documented: {documented}（其余为 demo_estimate，引擎将标注 indicative）")
    print(f"ige_materials: {len(ige)} 条")


if __name__ == "__main__":
    main()
